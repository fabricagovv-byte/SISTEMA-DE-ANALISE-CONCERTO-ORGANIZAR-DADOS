import io
from pathlib import Path

import pandas as pd

from agente import executar, gerar_excel
from agente import validadores as v
from agente.agente_ia import _mascarar, aplicar_mapeamento
from agente.leitor import ler_arquivo
from agente.normalizador import normalizar_tabela
from agente.sincronizador import sincronizar

EXEMPLOS = Path(__file__).resolve().parent.parent / "exemplos"


def test_cpf_cnpj():
    assert v.cpf_valido("529.982.247-25")
    assert not v.cpf_valido("529.982.247-24")
    assert not v.cpf_valido("111.111.111-11")
    assert v.cnpj_valido("11.222.333/0001-81")
    assert v.cnpj_valido("00000000000191")  # Banco do Brasil, zeros à esquerda
    assert not v.cnpj_valido("11.222.333/0001-80")
    assert v.formatar_cnpj("191") == "00.000.000/0001-91"


def test_valores_e_datas():
    assert v.converter_valor_monetario("R$ 1.234,56") == 1234.56
    assert v.converter_valor_monetario("1,234.56") == 1234.56
    assert v.converter_valor_monetario("(1.000,00)") == -1000.0
    assert v.converter_valor_monetario("abc") is None
    assert str(v.converter_data("05/03/2024")) == "2024-03-05"
    assert str(v.converter_data("2024-03-05T10:00:00.000")) == "2024-03-05"
    assert v.converter_data("31/02/2024") is None
    assert v.normalizar_uf("São Paulo") == "SP"
    assert v.formatar_cep("1310100") == "01310-100"


def test_csv_latin1_ponto_e_virgula_preserva_zeros():
    conteudo = "CPF;Município\n05298224725;São Paulo\n".encode("latin-1")
    df = ler_arquivo("dados.csv", conteudo)["dados"]
    assert df.loc[0, "CPF"] == "05298224725"
    assert df.loc[0, "Município"] == "São Paulo"


def test_normalizacao_detecta_tipos_e_problemas():
    df = pd.DataFrame({
        "CPF Titular": ["529.982.247-25", "52998224725", "123"],
        "Valor Pago": ["R$ 10,00", "R$ 10,00", "x"],
        "Data": ["01/01/2024", "01/01/2024", "2024-02-01"],
    })
    r = normalizar_tabela("t", df)
    assert r.tipos == {"cpf_titular": "cpf", "valor_pago": "valor", "data": "data"}
    assert r.duplicadas_removidas == 1
    assert {p["problema"] for p in r.problemas} == {"CPF inválido", "Número/valor não reconhecido"}


def test_sincronizacao_por_cpf_com_divergencia():
    a = normalizar_tabela("a", pd.DataFrame({"cpf": ["529.982.247-25", "111.444.777-35"], "nome": ["Ana", "Beto"]}))
    b = normalizar_tabela("b", pd.DataFrame({
        "NR_CPF": ["52998224725", "52998224725"], "nome": ["Ana Maria", "Ana Maria"], "valor": ["10,00", "5,50"],
    }))
    s = sincronizar([a, b])
    assert s.chave == "cpf"
    assert len(s.consolidado) == 2
    linha = s.consolidado[s.consolidado["cpf"] == "529.982.247-25"].iloc[0]
    assert linha["qtd_fontes"] == 2
    assert linha["valor__b"] == 15.5
    assert list(s.divergencias["campo"]) == ["nome"]


def test_pipeline_completo_com_exemplos():
    arquivos = [(p.name, p.read_bytes()) for p in sorted(EXEMPLOS.iterdir())]
    resultado = executar(arquivos)
    assert resultado.sincronizacao.chave == "cpf"
    assert len(resultado.problemas) > 0
    planilha = pd.read_excel(io.BytesIO(gerar_excel(resultado)), sheet_name=None)
    assert {"Resumo", "Consolidado", "Problemas", "Dicionario"} <= set(planilha)


def test_ia_mascara_dados_pessoais_e_aplica_mapeamento():
    assert "529" not in _mascarar("529.982.247-25")
    mapa = aplicar_mapeamento({"tabelas": [{
        "tabela": "x", "descricao": "d",
        "colunas": [{"original": "NOME_FAVORECIDO", "nome_padronizado": "nome", "tipo": "texto"}],
    }]})
    assert mapa["x"][0] == {"NOME_FAVORECIDO": "nome"}
    assert mapa["x"][1] == {"nome": "texto"}


# ---------- Documentos, continuação e LGPD ----------

from agente import gerar_pacote_zip  # noqa: E402
from agente.documentos import Documento, ler_documento, parece_texto_corrido  # noqa: E402
from agente.organizador_docs import classificar, encontrar_continuacao, fichar  # noqa: E402
from agente.sensiveis import anonimizar_texto, encontrar_em_texto  # noqa: E402


def _doc(nome, texto):
    return Documento(nome, texto, [texto])


def test_classificacao_usa_o_titulo():
    assert classificar("NOTA DE EMPENHO 2026NE000001\nReferente ao Contrato nº 3/2026")[0] == "Nota de empenho"
    assert classificar("3º TERMO ADITIVO AO CONTRATO Nº 3/2026")[0] == "Termo aditivo"
    assert classificar("EDITAL DE PREGÃO ELETRÔNICO Nº 1/2026")[0] == "Edital"


def test_continuacao_por_processo_e_lacunas():
    docs = {
        "edital": _doc("edital", "EDITAL Nº 1/2026\nProcesso SEI nº 23000.000001/2026-01\nfls. 1"),
        "contrato": _doc("contrato", "CONTRATO Nº 7/2026\nProcesso 23000.000001/2026-01\nfls. 5"),
        "aditivo3": _doc("aditivo3", "3º TERMO ADITIVO AO CONTRATO Nº 7/2026"),
        "solto": _doc("solto", "OFÍCIO Nº 9/2026\nAssunto: outro tema"),
    }
    fichas = [fichar(d) for d in docs.values()]
    dossies, vinculos, lacunas = encontrar_continuacao(fichas, docs)
    grupo = dossies[dossies["documento"] == "edital"]["dossie"].iloc[0]
    assert set(dossies[dossies["dossie"] == grupo]["documento"]) == {"edital", "contrato", "aditivo3"}
    assert dossies["dossie"].nunique() == 2
    texto_lacunas = " ".join(lacunas["lacuna"])
    assert "1º termo aditivo" in texto_lacunas and "2º termo aditivo" in texto_lacunas
    assert "folhas 2 a 4" in texto_lacunas


def test_partes_de_arquivo_e_parte_faltando():
    docs = {n: _doc(n, f"RELATÓRIO texto {n}") for n in ("rel_parte1", "rel_parte3")}
    _, vinculos, lacunas = encontrar_continuacao([fichar(d) for d in docs.values()], docs)
    assert len(vinculos) == 1
    assert "Falta a parte 2 do arquivo" in list(lacunas["lacuna"])


def test_dados_sensiveis_em_texto_e_anonimizacao():
    texto = (
        "Encaminhamos a Sra. Maria Souza, CPF 529.982.247-25, RG nº 12.345.678-9, "
        "e-mail maria@exemplo.com, com diagnóstico CID F32. Criança sob guarda."
    )
    achados = encontrar_em_texto(texto, "oficio")
    tipos = {a["tipo"] for a in achados}
    assert {"CPF", "RG", "E-mail", "Nome de pessoa", "Saúde: CID", "Criança/adolescente"} <= tipos
    assert all("529.982.247-25" not in a["valor_mascarado"] + a["contexto"] for a in achados)
    anonimo = anonimizar_texto(texto)
    for sensivel in ("529.982.247-25", "Maria Souza", "maria@exemplo.com", "12.345.678-9", "F32"):
        assert sensivel not in anonimo


def test_cnpj_nao_e_anonimizado():
    assert "11.222.333/0001-81" in anonimizar_texto("Contratada CNPJ 11.222.333/0001-81")


def test_leitura_docx_html_e_txt():
    import docx

    arquivo = docx.Document()
    arquivo.add_paragraph("CONTRATO Nº 1/2026")
    tabela = arquivo.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("cpf", "valor"), ("52998224725", "10,00"), ("11144477735", "5,00")]):
        tabela.cell(i, 0).text, tabela.cell(i, 1).text = a, b
    saida = io.BytesIO()
    arquivo.save(saida)
    doc = ler_documento("c.docx", saida.getvalue())
    assert "CONTRATO" in doc.texto and len(doc.tabelas) == 1

    html = ler_documento("p.html", "<h1>PORTARIA Nº 5/2026</h1><script>x</script>".encode())
    assert "PORTARIA" in html.texto and "x" not in html.texto
    assert parece_texto_corrido("Ofício\nTexto corrido do documento.".encode())
    assert not parece_texto_corrido("a;b\n1;2\n3;4".encode())


def test_pipeline_completo_com_documentos():
    arquivos = [(p.name, p.read_bytes()) for p in sorted(EXEMPLOS.iterdir())]
    r = executar(arquivos)
    assert len(r.documentos) >= 5
    assert "ALTO" in set(r.risco["risco"])
    assert len(r.lacunas) > 0
    assert len(r.cruzamento_docs_dados) > 0
    planilha = pd.read_excel(io.BytesIO(gerar_excel(r)), sheet_name=None)
    assert {"Documentos", "Dossies", "Lacunas", "Dados_sensiveis_LGPD", "Risco_LGPD"} <= set(planilha)
    # Nenhum CPF completo pode aparecer na aba de dados sensíveis.
    aba = planilha["Dados_sensiveis_LGPD"].astype(str)
    assert not aba.apply(lambda c: c.str.contains(r"\d{3}\.\d{3}\.\d{3}-\d{2}")).any().any()
    import zipfile

    with zipfile.ZipFile(io.BytesIO(gerar_pacote_zip(r))) as z:
        nomes = z.namelist()
        assert any(n.startswith("anonimizado/documentos/") for n in nomes)
        oficio = z.read("anonimizado/documentos/oficio_atendimento_social.txt").decode()
        assert "[CPF]" in oficio and "Francisca" not in oficio


def test_gateway_compativel(monkeypatch):
    from agente import agente_ia

    monkeypatch.setenv("AGENTE_API_URL", "iron-exemplo.fly.dev/v1")
    assert agente_ia.endereco_api() == "https://iron-exemplo.fly.dev"
    assert agente_ia.usa_gateway()
    monkeypatch.setenv("AGENTE_API_URL", "https://api.anthropic.com")
    assert not agente_ia.usa_gateway()
    assert agente_ia._extrair_json('Claro! ```json\n{"a": "x}", "b": [1, {"c": 2}]}\n``` fim') == '{"a": "x}", "b": [1, {"c": 2}]}'


def test_gateway_mapeia_colunas_e_nao_envia_dados_pessoais(monkeypatch):
    from types import SimpleNamespace

    from agente import agente_ia

    enviados = []

    class Stream:
        def __init__(self, **kw):
            enviados.append(kw)
            assert "betas" not in kw and "fallbacks" not in kw  # modo compatível
        def __enter__(self):
            return self
        def __exit__(self, *a):
            return False
        def get_final_message(self):
            texto = '{"tabelas": [{"tabela": "t", "descricao": "d", "colunas": [{"original": "NR_CPF", "nome_padronizado": "cpf", "tipo": "cpf"}]}]}'
            return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text=texto)])

    monkeypatch.setenv("AGENTE_API_URL", "https://gateway.exemplo/v1")
    monkeypatch.setattr(agente_ia, "_cliente", lambda: SimpleNamespace(messages=SimpleNamespace(stream=Stream)))
    df = pd.DataFrame({"NR_CPF": ["529.982.247-25"], "EMAIL": ["maria@exemplo.com"]})
    resultado = agente_ia.mapear_colunas({"t": df})
    assert resultado["tabelas"][0]["colunas"][0]["nome_padronizado"] == "cpf"
    conteudo = enviados[0]["messages"][0]["content"]
    assert "529.982.247-25" not in conteudo and "maria@" not in conteudo
