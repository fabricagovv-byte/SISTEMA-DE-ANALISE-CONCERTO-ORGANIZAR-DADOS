"""Leitura de documentos públicos (PDF, DOCX, TXT, HTML, MD) em texto + tabelas."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

EXTENSOES_DOCUMENTO = {".pdf", ".docx", ".html", ".htm", ".md", ".rtf", ".log", ".eml"}


@dataclass
class Documento:
    nome: str
    texto: str
    paginas: list[str] = field(default_factory=list)
    tabelas: list[pd.DataFrame] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


def _decodificar(conteudo: bytes) -> str:
    for codificacao in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return conteudo.decode("latin-1", errors="replace")


def _tabela_valida(linhas: list[list]) -> pd.DataFrame | None:
    linhas = [[("" if c is None else re.sub(r"\s+", " ", str(c)).strip()) for c in l] for l in linhas if l]
    linhas = [l for l in linhas if any(l)]
    if len(linhas) < 3 or len(linhas[0]) < 2:
        return None
    cabecalho = [c or f"coluna_{i + 1}" for i, c in enumerate(linhas[0])]
    corpo = [l + [""] * (len(cabecalho) - len(l)) for l in linhas[1:]]
    return pd.DataFrame([l[: len(cabecalho)] for l in corpo], columns=cabecalho)


def _ler_pdf(conteudo: bytes) -> Documento:
    import pdfplumber

    paginas, tabelas, avisos = [], [], []
    with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
        for pagina in pdf.pages:
            paginas.append(pagina.extract_text() or "")
            for bruta in pagina.extract_tables():
                df = _tabela_valida(bruta)
                if df is not None:
                    tabelas.append(df)
    if paginas and sum(len(p.strip()) for p in paginas) < 20 * len(paginas):
        avisos.append(
            "PDF parece digitalizado (imagem): pouco texto extraído. Use um PDF pesquisável ou passe OCR antes."
        )
    return Documento("", "\n\n".join(paginas), paginas, tabelas, avisos)


def _ler_docx(conteudo: bytes) -> Documento:
    import docx

    arquivo = docx.Document(io.BytesIO(conteudo))
    partes = [p.text for p in arquivo.paragraphs]
    tabelas = []
    for tabela in arquivo.tables:
        linhas = [[celula.text for celula in linha.cells] for linha in tabela.rows]
        df = _tabela_valida(linhas)
        if df is not None:
            tabelas.append(df)
        partes.extend(" | ".join(l) for l in linhas)
    texto = "\n".join(partes)
    return Documento("", texto, [texto], tabelas)


def _ler_html(conteudo: bytes) -> Documento:
    from bs4 import BeautifulSoup

    sopa = BeautifulSoup(_decodificar(conteudo), "html.parser")
    for lixo in sopa(["script", "style", "nav", "footer"]):
        lixo.decompose()
    tabelas = []
    for tabela in sopa.find_all("table"):
        linhas = [[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])] for tr in tabela.find_all("tr")]
        df = _tabela_valida(linhas)
        if df is not None:
            tabelas.append(df)
    texto = sopa.get_text("\n", strip=True)
    return Documento("", texto, [texto], tabelas)


def _ler_rtf(conteudo: bytes) -> Documento:
    texto = _decodificar(conteudo)
    texto = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1252"), texto)
    texto = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", texto)
    return Documento("", texto, [texto])


def _ler_eml(conteudo: bytes) -> Documento:
    from email import policy
    from email.parser import BytesParser

    msg = BytesParser(policy=policy.default).parsebytes(conteudo)
    cabecalho = "\n".join(f"{k}: {msg[k]}" for k in ("From", "To", "Cc", "Subject", "Date") if msg[k])
    corpo = msg.get_body(preferencelist=("plain", "html"))
    texto = corpo.get_content() if corpo is not None else ""
    if corpo is not None and corpo.get_content_type() == "text/html":
        from bs4 import BeautifulSoup
        texto = BeautifulSoup(texto, "html.parser").get_text("\n", strip=True)
    anexos = [a.get_filename() for a in msg.iter_attachments() if a.get_filename()]
    doc = Documento("", f"{cabecalho}\n\n{texto}", [f"{cabecalho}\n\n{texto}"])
    if anexos:
        doc.avisos.append("e-mail com anexo(s) não processado(s): " + ", ".join(anexos))
    return doc


def ler_documento(nome: str, conteudo: bytes) -> Documento:
    extensao = Path(nome).suffix.lower()
    if extensao == ".pdf":
        doc = _ler_pdf(conteudo)
    elif extensao == ".docx":
        doc = _ler_docx(conteudo)
    elif extensao in {".html", ".htm"}:
        doc = _ler_html(conteudo)
    elif extensao == ".rtf":
        doc = _ler_rtf(conteudo)
    elif extensao == ".eml":
        doc = _ler_eml(conteudo)
    else:
        texto = _decodificar(conteudo)
        doc = Documento("", texto, [texto])
    doc.nome = Path(nome).stem
    return doc


def parece_texto_corrido(conteudo: bytes) -> bool:
    """Para .txt: True se for documento em prosa, False se for tabela delimitada."""
    linhas = [l for l in _decodificar(conteudo[:20000]).splitlines() if l.strip()][:30]
    if len(linhas) < 2:
        return True
    for sep in (";", "\t", "|", ","):
        contagens = [l.count(sep) for l in linhas]
        if contagens[0] >= 1 and sum(c == contagens[0] for c in contagens) / len(contagens) > 0.8:
            # Vírgula em frases ("Fulano: precisa melhorar, afastado...") não é tabela:
            # tabela tem células curtas; prosa tem muitas palavras por "célula".
            celulas = [c for l in linhas for c in l.split(sep)]
            palavras = sum(len(c.split()) for c in celulas) / max(len(celulas), 1)
            if palavras > 3 or any(":" in l.split(sep)[0] for l in linhas[:3]):
                return True
            return False
    return True
